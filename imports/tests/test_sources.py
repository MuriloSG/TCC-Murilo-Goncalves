"""Capacidades do motor usadas pela importação de ensalamento, testadas sobre auth.Group.

A fonte de planilha em grade, os limites declarados pela importação, os campos próprios
do envio, os metadados da auditoria e a ausência de modelo são do motor, e por isso são
testados aqui, sem depender de nenhum app de domínio.
"""

import io

from django import forms
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook

from accounts.models import User
from audit import catalog
from audit.models import AuditLog
from imports.definitions import Column
from imports.forms import ImportUploadForm
from imports.models import Import
from imports.reader import FileRejected
from imports.tests.definitions import GroupImport
from imports.tests.worker import confirm, preview
from imports.workbook import WorkbookSource, cell_reference

SEMESTER_LABEL = 'Semestre'


class GridGroupImport(GroupImport):
    """Importação de grade: cada célula preenchida das abas declaradas vira um grupo."""

    kind = 'GRUPO_GRADE'
    columns = GroupImport.columns + (Column('Célula'), Column(SEMESTER_LABEL))
    identifier = 'Célula'
    sheets = ('Primeira', 'Segunda')
    has_template = False
    upload_fields = {SEMESTER_LABEL: forms.CharField(label=SEMESTER_LABEL, max_length=7)}

    def source(self):
        return WorkbookSource()

    def read_cells(self, workbook):
        for title in self.sheets:
            sheet = workbook[title]
            for row in range(1, sheet.max_row + 1):
                for column in range(1, sheet.max_column + 1):
                    value = sheet.cell(row, column).value
                    if value is None or not str(value).strip():
                        continue
                    yield {'Nome': str(value), 'Célula': cell_reference(title, row, column)}

    def audit_metadata(self, rows):
        codes = {row.values.get(SEMESTER_LABEL) for row in rows}
        return {'semestre': codes.pop() if len(codes) == 1 else None}


def workbook_file(sheets, name='planilha.xlsx'):
    book = Workbook()
    book.remove(book.active)
    for title, cells in sheets.items():
        sheet = book.create_sheet(title)
        for coordinate, value in cells.items():
            sheet[coordinate] = value
    stream = io.BytesIO()
    book.save(stream)
    return SimpleUploadedFile(
        name,
        stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def two_sheets():
    return {'Primeira': {'A1': 'Docentes', 'B2': 'Técnicos'}, 'Segunda': {'A1': 'Direção'}}


class WorkbookSourceTests(TestCase):
    def setUp(self):
        self.definition = GridGroupImport()

    def read(self, uploaded):
        return self.definition.source().read(uploaded, self.definition)

    def assertRejected(self, uploaded, message):
        with self.assertRaises(FileRejected) as context:
            self.read(uploaded)
        self.assertIn(message, context.exception.message)

    def test_reads_one_row_per_cell_of_the_declared_sheets(self):
        rows = self.read(workbook_file(two_sheets()))
        self.assertEqual([row.line_number for row in rows], [1, 2, 3])
        self.assertEqual(
            [row.values['Célula'] for row in rows], ['Primeira!A1', 'Primeira!B2', 'Segunda!A1'],
        )
        self.assertEqual([row.values['Nome'] for row in rows], ['Docentes', 'Técnicos', 'Direção'])
        self.assertEqual([row.errors for row in rows], [[], [], []])

    def test_extra_sheet_is_ignored(self):
        sheets = two_sheets() | {'Calculada': {'A1': 'não importa'}}
        rows = self.read(workbook_file(sheets))
        self.assertEqual([row.values['Nome'] for row in rows], ['Docentes', 'Técnicos', 'Direção'])

    def test_missing_declared_sheet_rejects_the_file(self):
        self.assertRejected(workbook_file({'Primeira': {'A1': 'Docentes'}}), 'Falta a aba "Segunda"')

    def test_file_that_is_not_a_workbook_is_rejected(self):
        uploaded = SimpleUploadedFile('planilha.xlsx', b'Nome;Perfil\n', content_type='text/csv')
        self.assertRejected(uploaded, 'não pôde ser lido como planilha')

    def test_workbook_without_data_is_rejected(self):
        self.assertRejected(workbook_file({'Primeira': {}, 'Segunda': {}}), 'não tem linhas de dados')

    def test_values_are_kept_raw(self):
        rows = self.read(workbook_file({'Primeira': {'A1': ' Docentes\n do campus '}, 'Segunda': {'A1': 'Direção'}}))
        self.assertEqual(rows[0].values['Nome'], ' Docentes\n do campus ')

    def test_formula_without_cached_value_is_an_empty_cell(self):
        sheets = {'Primeira': {'A1': '=UPPER(B1)', 'A2': 'Docentes'}, 'Segunda': {'A1': 'Direção'}}
        rows = self.read(workbook_file(sheets))
        self.assertEqual([row.values['Nome'] for row in rows], ['Docentes', 'Direção'])


class LimitTests(TestCase):
    def setUp(self):
        self.definition = GridGroupImport()

    @override_settings(IMPORT_MAX_ROWS=2)
    def test_definition_limit_replaces_the_configured_row_limit(self):
        self.definition.max_rows = 3
        rows = self.definition.source().read(workbook_file(two_sheets()), self.definition)
        self.assertEqual(len(rows), 3)

    @override_settings(IMPORT_MAX_ROWS=10)
    def test_definition_limit_refuses_above_its_own_value(self):
        self.definition.max_rows = 2
        with self.assertRaises(FileRejected) as context:
            self.definition.source().read(workbook_file(two_sheets()), self.definition)
        self.assertIn('O arquivo tem 3 linhas. O limite é 2.', context.exception.message)

    @override_settings(IMPORT_MAX_FILE_SIZE=10)
    def test_definition_size_limit_replaces_the_configured_one(self):
        self.definition.max_file_size = 1024 * 1024
        rows = self.definition.source().read(workbook_file(two_sheets()), self.definition)
        self.assertEqual(len(rows), 3)

    @override_settings(IMPORT_MAX_ROWS=2)
    def test_configured_limit_applies_without_a_declared_one(self):
        with self.assertRaises(FileRejected) as context:
            self.definition.source().read(workbook_file(two_sheets()), self.definition)
        self.assertIn('O limite é 2.', context.exception.message)


class UploadFieldTests(TestCase):
    def setUp(self):
        self.definition = GridGroupImport()

    def form(self, data=None, uploaded=None):
        return ImportUploadForm(
            data=data if data is not None else {'own_semestre': '2026/1'},
            files={'file': uploaded if uploaded is not None else workbook_file(two_sheets())},
            definition=self.definition,
        )

    def test_own_field_appears_beside_the_file(self):
        self.assertEqual(list(self.form().fields), ['file', 'own_semestre'])

    def test_own_value_is_recorded_in_every_row(self):
        form = self.form()
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual({row.values[SEMESTER_LABEL] for row in form.cleaned_data['rows']}, {'2026/1'})

    def test_missing_own_required_field_is_a_form_error(self):
        form = self.form(data={})
        self.assertFalse(form.is_valid())
        self.assertIn('own_semestre', form.errors)
        self.assertNotIn('rows', form.cleaned_data)

    def test_csv_import_keeps_its_own_field_texts(self):
        form = ImportUploadForm(data={}, files={}, definition=GroupImport())
        self.assertEqual(list(form.fields), ['file'])
        self.assertEqual(form.fields['file'].label, 'Arquivo CSV')
        self.assertFalse(form.is_valid())
        self.assertEqual(form.errors['file'], ['Escolha o arquivo CSV.'])


class AuditMetadataTests(TestCase):
    def setUp(self):
        self.executor = User.objects.create(email='admin@ifnmg.edu.br', name='Admin')
        self.definition = GridGroupImport()

    def preview(self):
        form = ImportUploadForm(
            data={'own_semestre': '2026/1'},
            files={'file': workbook_file(two_sheets())},
            definition=self.definition,
        )
        self.assertTrue(form.is_valid(), form.errors)
        return preview(
            self.definition, form.cleaned_data['rows'], form.cleaned_data['file'].name, self.executor,
        )

    def record(self, action):
        return AuditLog.objects.filter(action=action).latest('pk')

    def test_preview_record_has_the_definition_metadata(self):
        self.preview()
        metadata = self.record(catalog.CREATE).metadata
        self.assertEqual(metadata['semestre'], '2026/1')
        self.assertEqual((metadata['arquivo'], metadata['situacao'], metadata['linhas']), ('planilha.xlsx', 'VALIDADA', 3))

    def test_confirmation_record_is_the_summary_of_the_import(self):
        importation = self.preview()
        self.assertTrue(confirm(importation, self.definition, self.executor))
        record = self.record(catalog.IMPORT)
        self.assertEqual(importation.resolved_by, self.executor)
        self.assertEqual(str(record.resource_id), str(importation.pk))
        self.assertEqual(
            record.metadata,
            {
                'arquivo': 'planilha.xlsx',
                'situacao': 'CONCLUIDA',
                'cadastrados': 3,
                'atualizados': 0,
                'sem_alteracao': 0,
                'com_erro': 0,
                'semestre': '2026/1',
            },
        )
        self.assertIsNotNone(record.created_at)

    def test_only_the_import_itself_is_recorded(self):
        importation = self.preview()
        confirm(importation, self.definition, self.executor)
        self.assertEqual(Group.objects.count(), 3)
        self.assertEqual(
            list(AuditLog.objects.order_by('pk').values_list('action', 'resource')),
            [(catalog.CREATE, self.definition.audit_resource), (catalog.IMPORT, self.definition.audit_resource)],
        )

    def test_discard_record_has_the_definition_metadata(self):
        importation = self.preview()
        importation.discard(self.definition, self.executor)
        self.assertEqual(self.record(catalog.DISCARD).metadata['semestre'], '2026/1')
