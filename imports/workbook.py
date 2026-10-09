"""Leitura de XLSX em grade, com o pacote e as dimensões conferidos antes de percorrer as células."""

import zipfile
import zlib

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from imports.reader import (
    FileRejected,
    SourceRow,
    check_file_size,
    check_identifiers,
    check_row_count,
    listing,
    quantity,
)


def cell_reference(sheet_title, row, column):
    """Origem de uma célula no formato `Aba!Célula`, como a planilha a mostra."""
    return f'{sheet_title}!{get_column_letter(column)}{row}'


NOT_XLSX_MESSAGE = 'O arquivo não pôde ser lido como planilha. Envie a planilha no formato XLSX.'
REQUIRED_ENTRIES = ('[Content_Types].xml', 'xl/workbook.xml')
CHUNK_SIZE = 64 * 1024


def check_package(uploaded_file):
    """Recusa o que não é XLSX, pacote com entradas demais e conteúdo que se expande além do limite."""
    try:
        with zipfile.ZipFile(uploaded_file) as package:
            entries = package.infolist()
            names = {entry.filename for entry in entries}
            if any(name not in names for name in REQUIRED_ENTRIES):
                raise FileRejected(NOT_XLSX_MESSAGE)
            if len(entries) > settings.IMPORT_MAX_XLSX_ENTRIES:
                raise FileRejected(
                    f'A planilha tem {quantity(len(entries))} partes internas. O limite é '
                    f'{quantity(settings.IMPORT_MAX_XLSX_ENTRIES)}. Salve-a de novo no Excel ou no LibreOffice.'
                )
            expanded_size(package, entries)
    except (zipfile.BadZipFile, zipfile.LargeZipFile, zlib.error, NotImplementedError, RuntimeError, EOFError, OSError, ValueError):
        raise FileRejected(NOT_XLSX_MESSAGE) from None
    finally:
        uploaded_file.seek(0)


def expanded_size(package, entries):
    """Soma o que as entradas de fato descompactam, pois o tamanho declarado no ZIP pode mentir."""
    limit = settings.IMPORT_MAX_XLSX_UNCOMPRESSED_SIZE
    total = 0
    for entry in entries:
        with package.open(entry) as content:
            while chunk := content.read(CHUNK_SIZE):
                total += len(chunk)
                if total > limit:
                    raise FileRejected(
                        f'O conteúdo da planilha passa de {quantity(limit // (1024 * 1024))} MB quando '
                        'descompactado. Exclua abas, imagens ou linhas sem uso e envie de novo.'
                    )
    return total


def check_dimensions(workbook):
    """A leitura percorre as abas até `max_row` e `max_column`, criando cada célula que acessa."""
    rows, columns = settings.IMPORT_MAX_SHEET_ROWS, settings.IMPORT_MAX_SHEET_COLUMNS
    for sheet in workbook.worksheets:
        if sheet.max_row > rows or sheet.max_column > columns:
            raise FileRejected(
                f'A aba "{sheet.title}" vai até a linha {quantity(sheet.max_row)} e a coluna '
                f'{get_column_letter(min(sheet.max_column, 16384))}. O limite é a linha {quantity(rows)} e a '
                f'coluna {get_column_letter(columns)}. Exclua as linhas e colunas sem uso no fim da aba.'
            )


def open_workbook(uploaded_file):
    check_package(uploaded_file)
    try:
        workbook = load_workbook(uploaded_file, data_only=True)
    except Exception:
        raise FileRejected(NOT_XLSX_MESSAGE) from None
    check_dimensions(workbook)
    return workbook


class WorkbookSource:
    """Fonte de planilha em grade, para uma importação que declara abas e percorre células."""

    file_label = 'Arquivo da planilha'
    accept = '.xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    missing_message = 'Escolha o arquivo da planilha.'

    def help_text(self, definition):
        text = f'Planilha em XLSX, com as abas {listing(definition.sheets)}'
        optional = [f'"{name}"' for name in definition.optional_sheets]
        if optional:
            noun = 'a aba' if len(optional) == 1 else 'as abas'
            text += f' e, se houver, {noun} {listing(optional)}'
        return text + '.'

    def read(self, uploaded_file, definition):
        """Lê a planilha e devolve as linhas, uma por célula de dado das abas declaradas."""
        check_file_size(uploaded_file, definition)
        workbook = open_workbook(uploaded_file)
        missing = [f'"{name}"' for name in definition.sheets if name not in workbook.sheetnames]
        if missing:
            noun = 'a aba' if len(missing) == 1 else 'as abas'
            raise FileRejected(f'Falta {noun} {listing(missing)} na planilha. Confira o arquivo enviado.')
        rows = [SourceRow(number, values) for number, values in enumerate(definition.read_cells(workbook), start=1)]
        if not rows:
            raise FileRejected('O arquivo não tem linhas de dados.')
        check_row_count(rows, definition)
        check_identifiers(rows, definition)
        return rows
