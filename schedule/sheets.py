"""Leitura da estrutura das abas de horário da planilha institucional.

A grade de cada aba de horário tem blocos de dia empilhados: a coluna do dia traz o nome
mesclado no bloco, a coluna do horário traz a faixa ou as linhas de intervalo e de almoço,
e cada coluna de turma traz o texto da atividade. A linha 1 traz o nome do curso sobre o
grupo de colunas de uma mesma turma, e a primeira linha de cabeçalho traz o rótulo de cada
turma (D2, D5 e D6 da Change importar-ensalamento-da-planilha).

As abas calculadas da planilha, que repetem por fórmula o que está nestas, não são lidas:
`Ensalamento - Geral`, as abas de horário por docente e as abas por curso (D1).

O que impede a leitura da grade recusa o arquivo inteiro, com FileRejected. O que é do
conteúdo de uma célula vira erro da linha, em `schedule.activities`.
"""

import re
from dataclasses import dataclass
from datetime import time

from openpyxl.utils import get_column_letter

from imports.reader import FileRejected, listing, normalize_header
from schedule.models import Stage, Weekday
from schedule.utils import normalize_name

DAY_COLUMN = 1
SLOT_COLUMN = 2
FIRST_GROUP_COLUMN = 3
HEADER_MARK = 'DIA'
BREAKS = ('INTERVALO', 'ALMOÇO')
SLOT_PATTERN = re.compile(r'^(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})$')
LABEL_SEPARATOR = ' - '

WEEKDAYS = {
    'SEGUNDA': Weekday.MONDAY,
    'TERÇA': Weekday.TUESDAY,
    'TERCA': Weekday.TUESDAY,
    'QUARTA': Weekday.WEDNESDAY,
    'QUINTA': Weekday.THURSDAY,
    'SEXTA': Weekday.FRIDAY,
}

UNDERGRADUATE_SHEET = 'Horário - Graduação'
SECONDARY_SHEET = 'Horário - Ensino Médio'
CATALOG_SHEET = 'Validação de Dados'
TEACHER_CATALOG_COLUMN = 1
SPACE_CATALOG_COLUMN = 2
CATALOG_FIRST_ROW = 2
TEACHER_KIND = 'DOCENTE'
SPACE_KIND = 'SALA'
ACTIVITY_KIND = 'ATIVIDADE'
USERS_SHEET = 'usuários - sistema'
USER_KIND = 'USUARIO'
USER_EMAIL_COLUMN = 'E-mail'
USER_NAME_COLUMN = 'Nome'
USER_ROLE_COLUMN = 'Perfil'
USER_USERNAME_COLUMN = 'Nome de usuário'
USER_COLUMNS = (USER_EMAIL_COLUMN, USER_NAME_COLUMN, USER_ROLE_COLUMN, USER_USERNAME_COLUMN)
USER_HEADER_ROW = 1


@dataclass(frozen=True)
class SheetLayout:
    """Aba de horário a ser lida, com a etapa de ensino das turmas dela."""

    title: str
    stage: str


SHEET_LAYOUTS = (
    SheetLayout(UNDERGRADUATE_SHEET, Stage.UNDERGRADUATE),
    SheetLayout(SECONDARY_SHEET, Stage.SECONDARY),
)


@dataclass(frozen=True)
class GroupColumn:
    """Coluna de turma, com os textos crus do cabeçalho e do título do curso.

    Os textos ficam como estão na planilha porque são eles que vão para os valores da linha:
    a confirmação reexecuta a linha a partir deles, sem o arquivo, e precisa chegar ao mesmo
    rótulo, ao mesmo curso e ao mesmo aviso que a prévia (D2, D5, D16).
    """

    index: int
    header_text: str
    banner_text: str


def text(value):
    """Valor da célula como texto de uma linha, sem quebras e sem espaços repetidos."""
    return ' '.join(str(value if value is not None else '').replace('\n', ' ').split())


def raw_text(value):
    """Valor da célula como está na planilha, só convertido para texto."""
    return '' if value is None else str(value)


def header_rows(sheet):
    """Linhas de cabeçalho da aba, reconhecidas pela marca na coluna do dia."""
    return [row for row in range(1, sheet.max_row + 1) if text(sheet.cell(row, DAY_COLUMN).value) == HEADER_MARK]


def parse_slot(value):
    """Início e fim da faixa de horário, ou None quando a linha não é de faixa."""
    content = text(value)
    if not content or content.upper() in BREAKS:
        return None
    found = SLOT_PATTERN.match(content)
    if found is None:
        raise ValueError(content)
    start_hour, start_minute, end_hour, end_minute = (int(part) for part in found.groups())
    return time(start_hour, start_minute), time(end_hour, end_minute)


def parse_weekday(value):
    """Código do dia da semana, aceitando as grafias com e sem "-FEIRA"."""
    content = text(value).upper().removesuffix('-FEIRA').strip()
    return WEEKDAYS.get(content)


def banner_names(sheet, last_column):
    """Nome do curso da linha 1 por coluna, considerando as faixas mescladas."""
    names = {}
    for merged in sheet.merged_cells.ranges:
        if merged.min_row != 1:
            continue
        name = text(sheet.cell(1, merged.min_col).value)
        if not name:
            continue
        for column in range(merged.min_col, merged.max_col + 1):
            names[column] = name
    for column in range(FIRST_GROUP_COLUMN, last_column + 1):
        name = text(sheet.cell(1, column).value)
        if name and column not in names:
            names[column] = name
    return names


def split_header(value):
    """Rótulo da turma e nome curto do curso, a partir do cabeçalho da coluna.

    O cabeçalho tem duas linhas nas colunas em que o campus as usou, e uma só nas demais.
    Com uma linha, o rótulo é o cabeçalho inteiro e não há nome curto (D6).
    """
    lines = [normalize_name(line) for line in str(value or '').splitlines()]
    lines = [line for line in lines if line is not None]
    if not lines:
        return None, None
    label = lines[0].removeprefix(LABEL_SEPARATOR).removesuffix(LABEL_SEPARATOR)
    label = normalize_name(label.replace(LABEL_SEPARATOR, ' '))
    short_name = lines[1] if len(lines) > 1 else None
    return label, short_name


def group_columns(sheet, layout, first_header):
    """Colunas de turma da aba, com o rótulo e o curso de cada uma.

    O cabeçalho vem da primeira linha de cabeçalho, e não do cabeçalho repetido em cada
    bloco de dia, porque a planilha real tem coluna sem cabeçalho num bloco e grafias
    diferentes entre os blocos (D6).
    """
    last_column = sheet.max_column
    banners = banner_names(sheet, last_column)
    columns = {}
    without_label = []
    without_course = []
    for index in range(FIRST_GROUP_COLUMN, last_column + 1):
        header_text = raw_text(sheet.cell(first_header, index).value)
        label, short_name = split_header(header_text)
        if label is None:
            if column_has_activity(sheet, index):
                without_label.append(f'"{get_column_letter(index)}"')
            continue
        banner_text = banners.get(index, '')
        if not banner_text and short_name is None:
            without_course.append(f'"{get_column_letter(index)}"')
            continue
        columns[index] = GroupColumn(index=index, header_text=header_text, banner_text=banner_text)
    if without_label:
        noun = 'A coluna' if len(without_label) == 1 else 'As colunas'
        verb = 'tem atividade e não tem' if len(without_label) == 1 else 'têm atividade e não têm'
        raise FileRejected(
            f'{noun} {listing(without_label)} da aba "{layout.title}" {verb} cabeçalho de turma. '
            'Preencha o cabeçalho da coluna na planilha.'
        )
    if without_course:
        noun = 'A coluna' if len(without_course) == 1 else 'As colunas'
        raise FileRejected(
            f'{noun} {listing(without_course)} da aba "{layout.title}" não tem curso identificável, '
            'nem no título acima das colunas nem no cabeçalho da turma. Preencha o nome do curso na planilha.'
        )
    if not columns:
        raise FileRejected(f'A aba "{layout.title}" não tem nenhuma coluna de turma.')
    return columns


def column_has_activity(sheet, index):
    """Diz se a coluna tem alguma célula preenchida abaixo da primeira linha de cabeçalho."""
    marks = header_rows(sheet)
    if not marks:
        return False
    return any(text(sheet.cell(row, index).value) for row in range(marks[0] + 1, sheet.max_row + 1))


def day_blocks(sheet, layout):
    """Blocos de dia da aba: o dia da semana e o intervalo de linhas de cada um."""
    marks = header_rows(sheet)
    if len(marks) < 2:
        raise FileRejected(
            f'A aba "{layout.title}" não tem os blocos de dia esperados. '
            f'Confira a coluna com a palavra "{HEADER_MARK}".'
        )
    blocks = []
    unknown = []
    for position, mark in enumerate(marks[:-1]):
        weekday = parse_weekday(sheet.cell(mark + 1, DAY_COLUMN).value)
        if weekday is None:
            unknown.append(f'"{text(sheet.cell(mark + 1, DAY_COLUMN).value)}"')
            continue
        blocks.append((weekday, mark + 1, marks[position + 1] - 1))
    if unknown:
        raise FileRejected(
            f'A aba "{layout.title}" tem o dia da semana {listing(unknown)} fora dos dias esperados: '
            'segunda, terça, quarta, quinta e sexta.'
        )
    return blocks


def sheet_cells(sheet, layout):
    """Células de atividade da aba, com o dia, a faixa, a turma e o curso de cada uma."""
    marks = header_rows(sheet)
    columns = group_columns(sheet, layout, marks[0])
    for _, first_row, last_row in day_blocks(sheet, layout):
        for row in range(first_row, last_row + 1):
            try:
                slot = parse_slot(sheet.cell(row, SLOT_COLUMN).value)
            except ValueError as error:
                raise FileRejected(
                    f'A aba "{layout.title}" tem a faixa de horário "{error.args[0]}" na linha {row} '
                    'fora do formato 07:20 - 08:10.'
                ) from None
            if slot is None:
                continue
            for index, group in columns.items():
                content = raw_text(sheet.cell(row, index).value)
                if not text(content):
                    continue
                yield {
                    'day_text': raw_text(sheet.cell(first_row, DAY_COLUMN).value),
                    'slot_text': raw_text(sheet.cell(row, SLOT_COLUMN).value),
                    'group': group,
                    'content': content,
                    'reference': f'{layout.title}!{get_column_letter(index)}{row}',
                }


def catalog_entries(workbook):
    """Valores das listas de salas e de docentes, cada um com a célula de origem (D4).

    O valor vai cru para os valores da linha, e a retirada dos parênteses do nome da sala
    acontece na execução da linha, para que a confirmação chegue ao mesmo nome que a prévia.

    As salas vêm primeiro e os docentes depois, agrupados, para que a prévia mostre cada
    lista inteira de uma vez em vez de intercalar as duas.
    """
    sheet = workbook[CATALOG_SHEET]
    entries = []
    counts = {TEACHER_KIND: 0, SPACE_KIND: 0}
    for kind, column in ((SPACE_KIND, SPACE_CATALOG_COLUMN), (TEACHER_KIND, TEACHER_CATALOG_COLUMN)):
        for row in range(CATALOG_FIRST_ROW, sheet.max_row + 1):
            value = raw_text(sheet.cell(row, column).value)
            if not text(value):
                continue
            counts[kind] += 1
            entries.append({
                'kind': kind,
                'value': value,
                'reference': f'{CATALOG_SHEET}!{get_column_letter(column)}{row}',
            })
    empty = []
    if not counts[TEACHER_KIND]:
        empty.append('de docentes')
    if not counts[SPACE_KIND]:
        empty.append('de salas')
    if empty:
        raise FileRejected(
            f'A lista {listing(empty)} da aba "{CATALOG_SHEET}" está vazia. Preencha a lista na planilha.'
        )
    return entries


def strip_parentheses(value):
    """Retira os parênteses que envolvem o nome da sala na lista de salas.

    A lista real mistura `(1/10)` com `Lab. de Madeira`, e o nome da sala é o conteúdo (D4).
    """
    if value.startswith('(') and value.endswith(')'):
        return normalize_name(value[1:-1]) or value
    return value


def find_sheet(workbook, title):
    """Aba cujo nome é igual a `title` sem diferenciar caixa, acentos e bordas, ou None.

    Serve à aba opcional de usuários: como a ausência dela não gera aviso, um nome com
    outra caixa ou sem acento não pode passar por ausente (D4 da Change
    importar-usuarios-com-o-ensalamento).
    """
    wanted = normalize_header(title)
    return next((workbook[name] for name in workbook.sheetnames if normalize_header(name) == wanted), None)


def user_header_positions(sheet):
    """Coluna de cada rótulo do cabeçalho da aba de usuários, recusando coluna faltando ou repetida (D5).

    Coluna a mais é ignorada, como nas demais abas da grade.
    """
    labels = {normalize_header(label): label for label in USER_COLUMNS}
    positions, repeated = {}, []
    for column in range(1, sheet.max_column + 1):
        label = labels.get(normalize_header(text(sheet.cell(USER_HEADER_ROW, column).value)))
        if label is None:
            continue
        if label in positions:
            repeated.append(f'"{label}"')
        else:
            positions[label] = column
    expected = listing(USER_COLUMNS)
    if repeated:
        raise FileRejected(
            f'A coluna {listing(repeated)} aparece mais de uma vez na aba "{sheet.title}". Deixe só uma.'
        )
    missing = [f'"{label}"' for label in USER_COLUMNS if label not in positions]
    if missing:
        noun = 'a coluna' if len(missing) == 1 else 'as colunas'
        raise FileRejected(
            f'Falta {noun} {listing(missing)} na aba "{sheet.title}". Use as colunas {expected}.'
        )
    return positions


def user_entries(workbook):
    """Usuários da aba opcional, um por linha de dado, com a célula de origem (D4 e D5).

    Sem a aba, ou com a aba sem nenhuma célula preenchida, não há usuário e nada é
    recusado. A primeira linha é o cabeçalho, e as linhas vazias são ignoradas. Os valores
    vão crus, para que a confirmação normalize igual à prévia.
    """
    sheet = find_sheet(workbook, USERS_SHEET)
    if sheet is None:
        return []
    if not any(text(value) for row in sheet.iter_rows(values_only=True) for value in row):
        return []
    positions = user_header_positions(sheet)
    entries = []
    for row in range(USER_HEADER_ROW + 1, sheet.max_row + 1):
        values = {label: raw_text(sheet.cell(row, column).value) for label, column in positions.items()}
        if not any(text(value) for value in values.values()):
            continue
        entries.append({'values': values, 'reference': f'{sheet.title}!A{row}'})
    return entries
