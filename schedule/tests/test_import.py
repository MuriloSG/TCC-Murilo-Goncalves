"""Fluxo da importação de ensalamento pela máquina de estados do app `imports`.

A planilha destes testes é montada com a mesma estrutura da planilha real: linha de título
do curso, linha de cabeçalho por bloco de dia, coluna do dia, coluna do horário com as
linhas de intervalo e de almoço, colunas de turma e a aba de validação de dados com as duas
listas. O arquivo real é usado nos testes que conferem o comportamento diante dele.
"""

import io
from datetime import date, time
from unittest.mock import patch

from django.conf import settings
from django.db import DatabaseError
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook

from accounts import permissions
from accounts.models import Role, User
from audit import catalog
from audit.models import AuditLog
from imports.forms import ImportUploadForm
from imports.models import Import
from imports.tests.worker import confirm, preview, refresh
from outbox.models import OutboxEntry
from schedule.imports import RECURRENCE_CHANGED, SpaceAllocationImport, concluded_message
from schedule.models import (
    Allocation,
    ClassGroup,
    ClassOffering,
    Course,
    Origin,
    Semester,
    Space,
    Stage,
    Subject,
    Teacher,
    TeachingAssignment,
    TimeSlot,
    Weekday,
    teacher_of,
)
from schedule.sheets import CATALOG_SHEET, SECONDARY_SHEET, UNDERGRADUATE_SHEET, USER_COLUMNS, USERS_SHEET

SPREADSHEET = settings.BASE_DIR / 'sources' / 'planilha_example.xlsx'
SUGGESTED_SPREADSHEET = settings.BASE_DIR / 'sources' / 'planilha_example_sugerida.xlsx'
CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

TEACHERS = ['Viviane', 'Elza', 'Eduardo Garrido', 'Thaís Rios', 'Gera Jr.', 'Érica']
SPACES = ['(1/13)', '(2/6)', '(HV 3)', 'Lab. de Madeira']
UNDERGRADUATE_SLOTS = ['07:20 - 08:10', '08:10 - 09:00', 'INTERVALO', '09:20 - 10:10']
SECONDARY_SLOTS = ['07:20 - 08:10', 'ALMOÇO', '13:20 - 14:10']
SPACE_NAMES = ['1/13', '2/6', 'HV 3', 'Lab. de Madeira']


def grid(sheet, banner, headers, days, activities):
    """Monta uma aba de horário com blocos de dia, como na planilha real."""
    if banner:
        sheet['C1'] = banner
        sheet.merge_cells(start_row=1, start_column=3, end_row=1, end_column=2 + len(headers))
    row = 2
    for day, slots in days:
        sheet.cell(row, 1, 'DIA')
        sheet.cell(row, 2, 'HORÁRIO')
        for position, header in enumerate(headers):
            sheet.cell(row, 3 + position, header)
        row += 1
        for slot in slots:
            if slot == slots[0]:
                sheet.cell(row, 1, day)
            sheet.cell(row, 2, slot)
            for position in range(len(headers)):
                content = activities.get((day, slot, position))
                if content is not None:
                    sheet.cell(row, 3 + position, content)
            row += 1
    sheet.cell(row, 1, 'DIA')
    sheet.cell(row, 2, 'HORÁRIO')


def workbook_file(
    undergraduate=None, secondary=None, teachers=None, spaces=None, name='ensalamento.xlsx',
    users=None, users_title=USERS_SHEET, users_header=USER_COLUMNS,
):
    book = Workbook()
    book.remove(book.active)
    grid(
        book.create_sheet(UNDERGRADUATE_SHEET),
        'BACHARELADO EM ENGENHARIA DE TESTE',
        ['2º PERÍODO\nEng. Teste', '4º PERÍODO\nEng. Teste'],
        [('SEGUNDA ', UNDERGRADUATE_SLOTS), ('TERÇA', UNDERGRADUATE_SLOTS)],
        undergraduate if undergraduate is not None else {
            ('SEGUNDA ', '07:20 - 08:10', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13)',
            ('SEGUNDA ', '08:10 - 09:00', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13)',
            ('SEGUNDA ', '09:20 - 10:10', 1): 'Cálculo III (Viviane) (2/6)',
        },
    )
    grid(
        book.create_sheet(SECONDARY_SHEET),
        'AGROINDÚSTRIA',
        ['1º ANO A\nAgroindústria'],
        [('SEGUNDA-FEIRA', SECONDARY_SLOTS)],
        secondary if secondary is not None else {
            ('SEGUNDA-FEIRA', '13:20 - 14:10', 0): 'Biologia (Érica) (HV 3)',
        },
    )
    catalog_sheet = book.create_sheet(CATALOG_SHEET)
    catalog_sheet['A1'] = 'Disciplina (Professor)'
    catalog_sheet['B1'] = 'Salas Disponíveis'
    for position, teacher in enumerate(teachers if teachers is not None else TEACHERS, start=2):
        catalog_sheet.cell(position, 1, teacher)
    for position, space in enumerate(spaces if spaces is not None else SPACES, start=2):
        catalog_sheet.cell(position, 2, space)
    if users is not None:
        users_sheet = book.create_sheet(users_title)
        users_sheet.append(list(users_header))
        for user in users:
            users_sheet.append(list(user))
    stream = io.BytesIO()
    book.save(stream)
    return SimpleUploadedFile(name, stream.getvalue(), content_type=CONTENT_TYPE)


PERIOD = {'own_início': '2026-08-03', 'own_fim': '2026-12-18'}


def space_manager(email='admin@ifnmg.edu.br', extra_permissions=()):
    """Usuário que importa o ensalamento e pode alterar o cadastro de espaços."""
    code = email.partition('@')[0].upper()
    role = Role.objects.create_role(
        code=code,
        name=code,
        permissions=[permissions.SPACE_ALLOCATION_IMPORT, *permissions.SPACE_PERMISSIONS, *extra_permissions],
    )
    return User.objects.register(email=email, name=code.capitalize(), role=role)


class ImportTestCase(TestCase):
    def setUp(self):
        self.executor = space_manager()
        self.definition = SpaceAllocationImport()

    def preview(self, uploaded=None, semester='2026/1', period=None):
        form = ImportUploadForm(
            data={'own_semestre': semester, **(PERIOD if period is None else period)},
            files={'file': uploaded if uploaded is not None else workbook_file()},
            definition=self.definition,
        )
        self.assertTrue(form.is_valid(), form.errors)
        return preview(
            self.definition, form.cleaned_data['rows'], form.cleaned_data['file'].name, self.executor,
        )

    def conclude(self, uploaded=None, semester='2026/1', period=None):
        importation = self.preview(uploaded, semester, period)
        self.assertEqual(importation.status, Import.Status.VALIDATED, self.errors_of(importation))
        self.assertTrue(confirm(importation, self.definition, self.executor))
        return importation

    def errors_of(self, importation):
        return [
            (row.line_number, row.values.get('Célula'), row.errors)
            for row in importation.rows.all()
            if row.errors
        ]

    def allocation(self):
        return Allocation.objects.get(subject__name='QUÍMICA ORGÂNICA', time_slot__start_time='07:20')


class PreviewTests(ImportTestCase):
    def test_preview_plans_everything_and_changes_nothing(self):
        importation = self.preview()
        self.assertEqual(importation.status, Import.Status.VALIDATED)
        self.assertEqual(importation.rows.count(), len(TEACHERS) + len(SPACES) + 4)
        self.assertEqual(importation.summary()['create'], len(TEACHERS) + len(SPACE_NAMES) + 4)
        unused = importation.rows.get(values__Nome='Lab. de Madeira')
        self.assertEqual((unused.action, unused.warnings), ('create', []))
        for model in (Semester, Space, Teacher, Course, Subject, ClassGroup, TimeSlot, Allocation, ClassOffering, TeachingAssignment):
            self.assertFalse(model.objects.exists(), model.__name__)

    def test_catalog_row_keeps_the_raw_value_and_plans_the_clean_name(self):
        rows = list(self.preview().rows.order_by('line_number'))
        space_row = rows[0]
        self.assertEqual(space_row.values['Nome'], '(1/13)')
        self.assertEqual(space_row.values['Célula'], f'{CATALOG_SHEET}!B2')
        self.assertEqual(space_row.changes['name'], [None, '1/13'])
        teacher_row = next(row for row in rows if row.values['Tipo'] == 'DOCENTE')
        self.assertEqual(teacher_row.values['Nome'], 'Viviane')
        self.assertEqual(teacher_row.values['Célula'], f'{CATALOG_SHEET}!A2')

    def test_activity_row_keeps_the_values_of_the_spreadsheet(self):
        rows = list(self.preview().rows.order_by('line_number'))
        activity = next(row for row in rows if row.values['Célula'] == f'{UNDERGRADUATE_SHEET}!C3')
        self.assertEqual(activity.values['Dia'], 'SEGUNDA ')
        self.assertEqual(activity.values['Horário'], '07:20 - 08:10')
        self.assertEqual(activity.values['Turma'], '2º PERÍODO\nEng. Teste')
        self.assertEqual(activity.values['Curso'], 'BACHARELADO EM ENGENHARIA DE TESTE')
        self.assertEqual(activity.values['Atividade'], 'QUÍMICA ORGÂNICA (Elza) (1/13)')
        self.assertEqual(activity.changes, {'space': [None, '1/13'], 'teachers': [None, ['Elza']]})

    def test_file_without_the_catalog_sheet_is_refused(self):
        book = Workbook()
        book.remove(book.active)
        grid(book.create_sheet(UNDERGRADUATE_SHEET), 'CURSO', ['2º PERÍODO\nX'], [('SEGUNDA', ['07:20 - 08:10'])], {})
        grid(book.create_sheet(SECONDARY_SHEET), 'CURSO', ['1º ANO A\nX'], [('SEGUNDA', ['07:20 - 08:10'])], {})
        stream = io.BytesIO()
        book.save(stream)
        uploaded = SimpleUploadedFile('sem_catalogo.xlsx', stream.getvalue(), content_type=CONTENT_TYPE)
        form = ImportUploadForm(data={'own_semestre': '2026/1', **PERIOD}, files={'file': uploaded}, definition=self.definition)
        self.assertFalse(form.is_valid())
        self.assertIn(CATALOG_SHEET, form.errors['file'][0])
        self.assertFalse(Import.objects.exists())

    def test_empty_catalog_is_refused(self):
        form = ImportUploadForm(
            data={'own_semestre': '2026/1', **PERIOD},
            files={'file': workbook_file(teachers=[], spaces=[])},
            definition=self.definition,
        )
        self.assertFalse(form.is_valid())
        self.assertIn('está vazia', form.errors['file'][0])

    def test_semester_is_required_in_the_upload(self):
        form = ImportUploadForm(data={}, files={'file': workbook_file()}, definition=self.definition)
        self.assertFalse(form.is_valid())
        self.assertEqual(form.errors['own_semestre'], ['Informe o semestre da planilha.'])
        self.assertFalse(Import.objects.exists())

    def test_semester_outside_the_format_is_refused_in_the_upload(self):
        form = ImportUploadForm(
            data={'own_semestre': '2026'}, files={'file': workbook_file()}, definition=self.definition,
        )
        self.assertFalse(form.is_valid())
        self.assertEqual(form.errors['own_semestre'], ['Informe o semestre no formato 2026/1.'])

    def test_unknown_teacher_is_an_error_of_the_activity_row(self):
        uploaded = workbook_file(undergraduate={
            ('SEGUNDA ', '07:20 - 08:10', 0): 'Física (Agnaldo) (1/13)',
        })
        importation = self.preview(uploaded)
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        row = importation.rows.get(values__Célula=f'{UNDERGRADUATE_SHEET}!C3')
        self.assertEqual(row.errors[0]['column'], 'Atividade')
        self.assertEqual(row.errors[0]['code'], 'docente_desconhecido')
        self.assertFalse(Teacher.objects.filter(name='Agnaldo').exists())


class ConfirmationTests(ImportTestCase):
    def test_confirmation_creates_the_semester_and_the_entities(self):
        self.conclude()
        self.assertEqual(list(Semester.objects.values_list('code', flat=True)), ['2026/1'])
        self.assertEqual(Teacher.objects.count(), len(TEACHERS))
        self.assertEqual(sorted(Space.objects.values_list('name', flat=True)), SPACE_NAMES)
        self.assertEqual(set(Space.objects.values_list('origin', flat=True)), {'PLANILHA'})
        self.assertEqual(Course.objects.count(), 2)
        self.assertEqual(Subject.objects.count(), 3)
        self.assertEqual(ClassGroup.objects.count(), 3)
        self.assertEqual(Allocation.objects.count(), 4)

    def test_allocation_of_a_consecutive_block_is_not_merged(self):
        self.conclude()
        chemistry = Allocation.objects.filter(subject__name='QUÍMICA ORGÂNICA').order_by('time_slot__start_time')
        self.assertEqual(chemistry.count(), 2)
        self.assertEqual(
            [(str(item.time_slot), item.space.name) for item in chemistry],
            [('07:20 - 08:10', '1/13'), ('08:10 - 09:00', '1/13')],
        )

    def test_allocation_keeps_the_details_of_the_activity(self):
        self.conclude()
        allocation = self.allocation()
        self.assertEqual(allocation.semester.code, '2026/1')
        self.assertEqual(allocation.class_group.label, '2º PERÍODO')
        self.assertEqual(allocation.class_group.course.name, 'BACHARELADO EM ENGENHARIA DE TESTE')
        self.assertEqual(allocation.class_group.course.stage, Stage.UNDERGRADUATE)
        self.assertEqual(allocation.weekday, Weekday.MONDAY)
        self.assertEqual(allocation.space.name, '1/13')
        self.assertEqual([teacher.name for teacher in allocation.teachers.all()], ['Elza'])
        self.assertEqual(allocation.origin, Origin.IMPORTED)

    def test_secondary_sheet_gives_its_own_stage(self):
        self.conclude()
        group = ClassGroup.objects.get(label='1º ANO A')
        self.assertEqual(group.course.name, 'AGROINDÚSTRIA')
        self.assertEqual(group.course.stage, Stage.SECONDARY)
        self.assertEqual(group.allocations.count(), 1)

    def test_space_and_teacher_without_activity_are_registered(self):
        self.conclude()
        space = Space.objects.get(name='Lab. de Madeira')
        self.assertEqual((space.is_active, space.allocations.count()), (True, 0))
        self.assertTrue(Teacher.objects.filter(name='Gera Jr.').exists())

    def test_repeated_catalog_value_becomes_one_record(self):
        uploaded = workbook_file(teachers=['Viviane', 'Elza', 'Érica', 'viviane'], spaces=['(1/13)', '(2/6)', '(HV 3)'])
        importation = self.conclude(uploaded)
        self.assertEqual(Teacher.objects.filter(name__iexact='viviane').count(), 1)
        self.assertEqual(Teacher.objects.count(), 3)
        repeated = importation.rows.get(values__Nome='viviane')
        first = importation.rows.get(values__Nome='Viviane')
        self.assertEqual((first.action, repeated.action), ('create', 'unchanged'))
        self.assertIsNone(repeated.target)

    def test_existing_space_cited_only_in_the_cells_is_recognized(self):
        self.conclude()
        fewer = workbook_file(
            spaces=['(1/13)', '(2/6)'],
            secondary={('SEGUNDA-FEIRA', '13:20 - 14:10', 0): 'Biologia (Érica) (HV 3)'},
        )
        importation = self.preview(fewer)
        self.assertEqual(importation.status, Import.Status.VALIDATED)
        row = importation.rows.get(values__Célula=f'{SECONDARY_SHEET}!C5')
        self.assertEqual((row.errors, row.pending, row.action), ([], [], 'unchanged'))

    def test_concluded_message_has_the_totals(self):
        importation = self.conclude()
        self.assertEqual(
            concluded_message(importation.summary(), importation),
            f'Importação do ensalamento concluída. {len(TEACHERS) + len(SPACE_NAMES) + 4} registros cadastrados, '
            '4 espaços cadastrados.',
        )


class ReimportTests(ImportTestCase):
    def test_reimport_without_changes_leaves_everything_unchanged(self):
        self.conclude()
        counts = {model: model.objects.count() for model in (Space, Teacher, Course, Subject, ClassGroup, Allocation)}
        importation = self.conclude()
        summary = importation.summary()
        self.assertEqual(summary['create'], 0)
        self.assertEqual(summary['update'], 0)
        self.assertEqual(summary['unchanged'], importation.rows.count())
        for model, count in counts.items():
            self.assertEqual(model.objects.count(), count, model.__name__)

    def test_allocation_that_changed_space_is_updated(self):
        self.conclude()
        changed = workbook_file(undergraduate={
            ('SEGUNDA ', '07:20 - 08:10', 0): 'QUÍMICA ORGÂNICA (Elza) (2/6)',
            ('SEGUNDA ', '08:10 - 09:00', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13)',
            ('SEGUNDA ', '09:20 - 10:10', 1): 'Cálculo III (Viviane) (2/6)',
        })
        importation = self.preview(changed)
        row = importation.rows.get(values__Célula=f'{UNDERGRADUATE_SHEET}!C3')
        self.assertEqual((row.action, row.changes), ('update', {'space': ['1/13', '2/6']}))
        self.assertEqual(self.allocation().space.name, '1/13')
        self.assertTrue(confirm(importation, self.definition, self.executor))
        self.assertEqual(self.allocation().space.name, '2/6')
        self.assertEqual(Allocation.objects.count(), 4)

    def test_allocation_that_changed_teachers_is_updated(self):
        self.conclude()
        changed = workbook_file(undergraduate={
            ('SEGUNDA ', '07:20 - 08:10', 0): 'QUÍMICA ORGÂNICA (Elza/ Viviane) (1/13)',
            ('SEGUNDA ', '08:10 - 09:00', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13)',
            ('SEGUNDA ', '09:20 - 10:10', 1): 'Cálculo III (Viviane) (2/6)',
        })
        importation = self.preview(changed)
        row = importation.rows.get(values__Célula=f'{UNDERGRADUATE_SHEET}!C3')
        self.assertEqual((row.action, row.changes), ('update', {'teachers': [['Elza'], ['Elza', 'Viviane']]}))
        confirm(importation, self.definition, self.executor)
        self.assertEqual(
            sorted(teacher.name for teacher in self.allocation().teachers.all()), ['Elza', 'Viviane'],
        )

    def test_allocation_missing_from_the_new_file_is_preserved(self):
        self.conclude()
        without_calculus = workbook_file(undergraduate={
            ('SEGUNDA ', '07:20 - 08:10', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13)',
            ('SEGUNDA ', '08:10 - 09:00', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13)',
        })
        importation = self.conclude(without_calculus)
        self.assertEqual(Allocation.objects.count(), 4)
        self.assertTrue(Allocation.objects.filter(subject__name='Cálculo III').exists())
        self.assertEqual([row.action for row in importation.rows.all() if row.action == 'delete'], [])

    def test_new_activity_is_added_to_a_semester_already_imported(self):
        self.conclude()
        with_more = workbook_file(undergraduate={
            ('SEGUNDA ', '07:20 - 08:10', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13)',
            ('SEGUNDA ', '08:10 - 09:00', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13)',
            ('SEGUNDA ', '09:20 - 10:10', 1): 'Cálculo III (Viviane) (2/6)',
            ('TERÇA', '07:20 - 08:10', 0): 'Biologia (Érica) (1/13)',
        })
        importation = self.conclude(with_more)
        self.assertEqual(importation.summary()['create'], 1)
        self.assertEqual(Allocation.objects.count(), 5)

    def test_manual_allocation_is_never_touched(self):
        self.conclude()
        semester = Semester.objects.get(code='2026/1')
        group = ClassGroup.objects.get(label='2º PERÍODO')
        subject, _ = Subject.objects.resolve('Reunião')
        slot = TimeSlot.objects.get(start_time='07:20')
        manual = Allocation.objects.create(
            semester=semester,
            class_group=group,
            subject=subject,
            weekday=Weekday.FRIDAY,
            time_slot=slot,
            origin=Origin.MANUAL,
        )
        self.conclude()
        manual.refresh_from_db()
        self.assertEqual(manual.origin, Origin.MANUAL)
        self.assertTrue(Allocation.objects.filter(pk=manual.pk).exists())

    def test_space_missing_from_the_new_file_is_preserved_and_active(self):
        self.conclude()
        fewer = workbook_file(
            spaces=['(1/13)', '(2/6)'],
            secondary={('SEGUNDA-FEIRA', '13:20 - 14:10', 0): 'Biologia (Érica) (1/13)'},
        )
        self.conclude(fewer)
        self.assertTrue(Space.objects.get(name='HV 3').is_active)

    def test_another_semester_is_untouched(self):
        self.conclude()
        first = sorted(Allocation.objects.filter(semester__code='2026/1').values_list('pk', flat=True))
        self.conclude(semester='2026/2')
        self.assertEqual(
            sorted(Allocation.objects.filter(semester__code='2026/1').values_list('pk', flat=True)), first,
        )
        self.assertEqual(Allocation.objects.filter(semester__code='2026/2').count(), 4)
        self.assertEqual(TimeSlot.objects.filter(semester__code='2026/2').count(), 4)

    def test_lesson_changed_in_the_system_is_kept_and_its_row_warns(self):
        self.conclude()
        chemistry = Allocation.objects.filter(subject__name='QUÍMICA ORGÂNICA').order_by('time_slot__start_time')
        superadmin = User.objects.create_superadmin('super@ifnmg.edu.br', 'Super', 'Senha-forte-123')
        with patch('django.utils.timezone.localdate', return_value=date(2026, 9, 1)):
            Allocation.objects.change_recurrence(
                [allocation.pk for allocation in chemistry], date(2026, 10, 5), superadmin,
                Weekday.MONDAY, time(7, 20), Space.objects.get(name='2/6'), [Teacher.objects.get(name='Elza')],
            )
        versions = list(Allocation.objects.order_by('pk').values())
        importation = self.preview()
        row = importation.rows.get(values__Célula=f'{UNDERGRADUATE_SHEET}!C3')
        self.assertEqual(row.action, 'unchanged')
        self.assertEqual(row.warnings, [RECURRENCE_CHANGED])
        self.assertTrue(confirm(importation, self.definition, self.executor))
        self.assertEqual(list(Allocation.objects.order_by('pk').values()), versions)

    def test_the_same_activity_twice_is_a_duplicate_error(self):
        uploaded = workbook_file(undergraduate={
            ('SEGUNDA ', '07:20 - 08:10', 0): 'QUÍMICA ORGÂNICA (Elza) (1/13) / QUÍMICA ORGÂNICA (Elza) (2/6)',
        })
        importation = self.preview(uploaded)
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        errors = [row.errors for row in importation.rows.all() if row.errors]
        self.assertEqual(len(errors), 2)
        self.assertEqual({found[0]['code'] for found in errors}, {'duplicate'})
        self.assertIn(f'{UNDERGRADUATE_SHEET}!C3#1', errors[0][0]['message'])
        self.assertIn(f'{UNDERGRADUATE_SHEET}!C3#2', errors[0][0]['message'])
        self.assertEqual(Allocation.objects.count(), 0)


class DivergenceTests(ImportTestCase):
    def test_data_changed_after_the_preview_gives_a_new_plan(self):
        importation = self.preview()
        Space.objects.register('1/13')
        self.assertFalse(confirm(importation, self.definition, self.executor))
        importation.refresh_from_db()
        self.assertEqual(importation.status, Import.Status.VALIDATED)
        self.assertEqual(importation.revalidation_count, 1)
        self.assertEqual(Allocation.objects.count(), 0)
        self.assertEqual((ClassOffering.objects.count(), TeachingAssignment.objects.count()), (0, 0))

    def test_confirmation_uses_the_recorded_semester(self):
        importation = self.preview(semester='2027/2')
        self.assertTrue(confirm(importation, self.definition, self.executor))
        self.assertEqual(list(Semester.objects.values_list('code', flat=True)), ['2027/2'])
        self.assertEqual(Allocation.objects.filter(semester__code='2027/2').count(), 4)


BLANK = {'own_início': '', 'own_fim': ''}
LATER = {'own_início': '2026-08-10', 'own_fim': '2026-12-18'}


class PeriodImportTests(ImportTestCase):
    def period(self, code='2026/1'):
        semester = Semester.objects.get(code=code)
        return semester.start_date, semester.end_date

    def last_summary(self):
        return AuditLog.objects.filter(action=catalog.IMPORT).latest('pk').metadata

    def form(self, period, semester='2026/1'):
        return ImportUploadForm(data={'own_semestre': semester, **period}, files={'file': workbook_file()}, definition=self.definition)

    def test_first_import_sets_the_period(self):
        self.conclude()
        self.assertEqual(self.period(), (date(2026, 8, 3), date(2026, 12, 18)))
        summary = self.last_summary()
        self.assertEqual((summary['inicio'], summary['fim']), ('2026-08-03', '2026-12-18'))
        self.assertNotIn('inicio_anterior', summary)

    def test_preview_and_discard_do_not_change_the_period(self):
        importation = self.preview()
        self.assertFalse(Semester.objects.filter(code='2026/1').exists())
        self.conclude()
        importation = self.preview(period=LATER)
        importation.discard(self.definition, self.executor)
        self.assertEqual(self.period(), (date(2026, 8, 3), date(2026, 12, 18)))

    def test_divergent_confirmation_does_not_change_the_period(self):
        self.conclude()
        importation = self.preview(period=LATER, uploaded=workbook_file(spaces=[*SPACES, 'Zoo II']))
        Space.objects.register('Zoo II')
        self.assertFalse(confirm(importation, self.definition, self.executor))
        self.assertEqual(self.period(), (date(2026, 8, 3), date(2026, 12, 18)))

    def test_reimport_with_new_dates_updates_and_audits_the_previous(self):
        self.conclude()
        self.conclude(period=LATER)
        self.assertEqual(self.period(), (date(2026, 8, 10), date(2026, 12, 18)))
        summary = self.last_summary()
        self.assertEqual(
            {key: summary[key] for key in ('inicio', 'fim', 'inicio_anterior', 'fim_anterior')},
            {'inicio': '2026-08-10', 'fim': '2026-12-18', 'inicio_anterior': '2026-08-03', 'fim_anterior': '2026-12-18'},
        )

    def test_reimport_with_blank_dates_keeps_the_period(self):
        self.conclude()
        self.conclude(period=BLANK)
        self.assertEqual(self.period(), (date(2026, 8, 3), date(2026, 12, 18)))
        summary = self.last_summary()
        self.assertEqual((summary['inicio'], summary['fim']), (None, None))
        self.assertNotIn('inicio_anterior', summary)

    def test_first_import_requires_the_dates(self):
        form = self.form(BLANK)
        self.assertFalse(form.is_valid())
        self.assertIn('Informe o início e o fim do semestre, que ainda não tem período.', form.errors['own_início'])
        self.assertIn('own_fim', form.errors)
        self.assertFalse(Semester.objects.exists())

    def test_summary_shows_the_dates_formatted(self):
        self.assertEqual(self.definition.upload_display('Início', '2026-08-03'), '03/08/2026')
        self.assertEqual(self.definition.upload_display('Fim', ''), 'Mantido')
        self.assertEqual(self.definition.upload_display('Semestre', '2026/1'), '2026/1')


class AuditTests(ImportTestCase):
    def test_summary_of_the_confirmation(self):
        importation = self.conclude()
        record = AuditLog.objects.filter(action=catalog.IMPORT).latest('pk')
        self.assertEqual(record.resource, catalog.SPACE_ALLOCATION_IMPORT)
        self.assertEqual(str(record.resource_id), str(importation.pk))
        self.assertEqual(
            record.metadata,
            {
                'arquivo': 'ensalamento.xlsx',
                'situacao': 'CONCLUIDA',
                'cadastrados': len(TEACHERS) + len(SPACE_NAMES) + 4,
                'atualizados': 0,
                'sem_alteracao': 0,
                'com_erro': 0,
                'semestre': '2026/1',
                'inicio': '2026-08-03',
                'fim': '2026-12-18',
                'espacos_cadastrados': 4,
                'outros_nomes_associados': 0,
                'espacos_reativados': 0,
            },
        )

    def test_no_record_for_each_entity_except_the_spaces(self):
        before = AuditLog.objects.latest('pk').pk
        self.conclude()
        self.assertEqual(
            sorted(set(AuditLog.objects.filter(pk__gt=before).values_list('resource', flat=True))),
            sorted([catalog.SPACE_ALLOCATION_IMPORT, catalog.SPACE]),
        )
        self.assertEqual(
            AuditLog.objects.filter(resource=catalog.SPACE, action=catalog.CREATE).count(), Space.objects.count(),
        )

    def test_preview_records_only_the_preview(self):
        before = AuditLog.objects.latest('pk').pk
        self.preview()
        self.assertEqual(
            list(AuditLog.objects.filter(pk__gt=before).values_list('action', 'resource')),
            [(catalog.CREATE, catalog.SPACE_ALLOCATION_IMPORT)],
        )

    def test_summary_of_a_reimport_counts_the_unchanged(self):
        self.conclude()
        importation = self.conclude()
        record = AuditLog.objects.filter(action=catalog.IMPORT).latest('pk')
        self.assertEqual(record.metadata['sem_alteracao'], importation.rows.count())
        self.assertEqual(record.metadata['cadastrados'], 0)
        self.assertEqual(record.metadata['espacos_cadastrados'], 0)


class RealSpreadsheetTests(ImportTestCase):
    def uploaded(self):
        return SimpleUploadedFile(
            'planilha_example.xlsx', SPREADSHEET.read_bytes(), content_type=CONTENT_TYPE,
        )

    def test_real_file_is_previewed_with_the_activities_that_the_campus_needs_to_fix(self):
        importation = self.preview(self.uploaded())
        self.assertEqual(importation.rows.count(), 1706)
        self.assertEqual(importation.file_name, 'planilha_example.xlsx')
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        summary = importation.summary()
        self.assertEqual(summary['errors'], 121)
        without_action = importation.rows.filter(action=None, errors=[], pending=[]).count()
        self.assertEqual(summary['create'], 1706 - 121 - summary['pending'] - summary['unchanged'] - without_action)
        self.assertEqual(summary['update'], 0)
        with self.assertRaises(Exception):
            confirm(importation, self.definition, self.executor)
        self.assertFalse(Allocation.objects.exists())


def offering_teachers(subject='QUÍMICA ORGÂNICA', label='2º PERÍODO', semester='2026/1'):
    offering = ClassOffering.objects.get(semester__code=semester, subject__name=subject, class_group__label=label)
    return {assignment.teacher.name: assignment.is_current for assignment in offering.assignments.select_related('teacher')}


def chemistry_by(teacher):
    return workbook_file(undergraduate={
        ('SEGUNDA ', '07:20 - 08:10', 0): f'QUÍMICA ORGÂNICA ({teacher}) (1/13)',
        ('SEGUNDA ', '08:10 - 09:00', 0): f'QUÍMICA ORGÂNICA ({teacher}) (1/13)',
        ('SEGUNDA ', '09:20 - 10:10', 1): 'Cálculo III (Viviane) (2/6)',
    })


class AcademicLinkTests(ImportTestCase):
    """Ofertas e docentes das ofertas estabelecidos pela importação (Change
    relacionar-professor-disciplina-turma)."""

    def test_confirmation_records_offerings_and_teachers(self):
        self.conclude()
        self.assertEqual(ClassOffering.objects.count(), 3)
        self.assertEqual(offering_teachers(), {'Elza': True})
        self.assertEqual(offering_teachers('Cálculo III', '4º PERÍODO'), {'Viviane': True})

    def test_refused_row_leaves_no_link(self):
        importation = self.preview(workbook_file(undergraduate={
            ('SEGUNDA ', '07:20 - 08:10', 0): 'QUÍMICA ORGÂNICA (Desconhecido) (1/13)',
        }))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        self.assertFalse(ClassOffering.objects.filter(subject__name='QUÍMICA ORGÂNICA').exists())

    def test_reimport_does_not_duplicate(self):
        self.conclude()
        counts = (ClassOffering.objects.count(), TeachingAssignment.objects.count())
        self.conclude()
        self.assertEqual((ClassOffering.objects.count(), TeachingAssignment.objects.count()), counts)

    def test_replaced_teacher_is_kept_as_history(self):
        self.conclude()
        self.conclude(chemistry_by('Viviane'))
        self.assertEqual(offering_teachers(), {'Elza': False, 'Viviane': True})

    def test_teacher_that_comes_back_is_current_again(self):
        self.conclude()
        self.conclude(chemistry_by('Viviane'))
        assignment = TeachingAssignment.objects.get(teacher__name='Elza')
        self.conclude()
        self.assertEqual(offering_teachers(), {'Elza': True, 'Viviane': False})
        self.assertTrue(TeachingAssignment.objects.get(pk=assignment.pk).is_current)

    def test_previous_semester_is_untouched(self):
        self.conclude()
        before = sorted(TeachingAssignment.objects.filter(offering__semester__code='2026/1').values_list('pk', 'is_current'))
        self.conclude(chemistry_by('Viviane'), semester='2026/2')
        after = sorted(TeachingAssignment.objects.filter(offering__semester__code='2026/1').values_list('pk', 'is_current'))
        self.assertEqual(after, before)
        self.assertEqual(offering_teachers(semester='2026/2'), {'Viviane': True})


ELZA = ('elza@ifnmg.edu.br', 'Elza Souza', 'DOCENTE', 'Elza')
VIVIANE = ('viviane@ifnmg.edu.br', 'Viviane Lima', 'DOCENTE', 'Viviane')
ERICA = ('erica@ifnmg.edu.br', 'Érica Alves', 'DOCENTE', 'Érica')


class UserSheetTests(ImportTestCase):
    """Aba opcional `usuários - sistema` (Change importar-usuarios-com-o-ensalamento)."""

    def setUp(self):
        self.executor = space_manager(extra_permissions=[permissions.USER_CREATE])
        self.definition = SpaceAllocationImport()
        self.teacher_role = Role.objects.get(code='DOCENTE')

    def with_users(self, *users, **options):
        return workbook_file(users=list(users), **options)

    def user_rows(self, importation):
        return list(importation.rows.filter(values__Tipo='USUARIO').order_by('line_number'))

    def imported_users(self):
        return User.objects.exclude(pk=self.executor.pk)

    def password_links(self):
        return OutboxEntry.objects.filter(kind='LINK_SENHA').exclude(reference=str(self.executor.pk))

    def user_audit(self):
        return AuditLog.objects.filter(resource=catalog.USER, action=catalog.CREATE).exclude(
            resource_id=str(self.executor.pk),
        )

    def error_codes(self, row):
        return [(error['column'], error['code']) for error in row.errors]

    def test_user_rows_are_recognized(self):
        importation = self.preview(self.with_users(ELZA, VIVIANE))
        rows = self.user_rows(importation)
        self.assertEqual([row.line_number for row in rows], [1, 2])
        self.assertEqual(rows[0].values['Célula'], 'usuários - sistema!A2')
        self.assertEqual(rows[0].values['E-mail'], 'elza@ifnmg.edu.br')
        self.assertIn('usuarios', [item.slug for item in self.definition.row_filters])

    def test_valid_users_are_planned(self):
        importation = self.preview(self.with_users(ELZA))
        self.assertEqual(importation.status, Import.Status.VALIDATED, self.errors_of(importation))
        row = self.user_rows(importation)[0]
        self.assertEqual(row.action, 'create')
        self.assertEqual(row.changes, {
            'email': [None, 'elza@ifnmg.edu.br'],
            'name': [None, 'Elza Souza'],
            'role': [None, 'DOCENTE'],
            'username': [None, 'Elza'],
        })

    def test_confirmation_creates_missing_users(self):
        importation = self.conclude(self.with_users(ELZA))
        user = User.objects.get(email='elza@ifnmg.edu.br')
        self.assertTrue(user.is_active)
        self.assertFalse(user.is_superadmin)
        self.assertEqual(user.role, self.teacher_role)
        self.assertEqual(user.username, 'Elza')
        self.assertFalse(user.has_usable_password())
        self.assertTrue(OutboxEntry.objects.filter(kind='LINK_SENHA', reference=str(user.pk)).exists())
        self.assertEqual(self.user_rows(importation)[0].created_object, str(user.pk))
        record = self.user_audit().get()
        self.assertEqual(record.origin, catalog.IMPORTED)
        self.assertTrue(Allocation.objects.exists())

    def test_role_is_applied(self):
        Role.objects.create_role(code='SECRETARIA', name='Secretaria')
        self.conclude(self.with_users(ELZA, ('ana@ifnmg.edu.br', 'Ana', 'secretaria', '')))
        self.assertEqual(User.objects.get(email='elza@ifnmg.edu.br').role.code, 'DOCENTE')
        ana = User.objects.get(email='ana@ifnmg.edu.br')
        self.assertEqual(ana.role.code, 'SECRETARIA')
        self.assertIsNone(ana.username)

    def test_created_user_matches_teacher(self):
        self.conclude(self.with_users(ELZA))
        user = User.objects.get(email='elza@ifnmg.edu.br')
        self.assertEqual(teacher_of(user), Teacher.objects.get(name='Elza'))

    def test_existing_user_is_not_duplicated(self):
        existing = User.objects.register(
            email='elza@ifnmg.edu.br', name='Elza Souza', role=self.teacher_role, username='Elza',
        )
        links = self.password_links().count()
        importation = self.conclude(self.with_users(('ELZA@ifnmg.edu.br', 'Elza Souza', 'DOCENTE', 'Elza')))
        row = self.user_rows(importation)[0]
        self.assertEqual((row.action, row.target, row.warnings), ('unchanged', str(existing.pk), []))
        self.assertEqual(User.objects.filter(email__iexact='elza@ifnmg.edu.br').count(), 1)
        self.assertEqual(self.password_links().count(), links)

    def test_existing_user_with_other_data_is_not_updated(self):
        Role.objects.create_role(code='SECRETARIA', name='Secretaria')
        existing = User.objects.register(
            email='elza@ifnmg.edu.br', name='Elza Souza', role=self.teacher_role, username='Elza',
        )
        importation = self.conclude(self.with_users(('elza@ifnmg.edu.br', 'Elza', 'SECRETARIA', 'Elza')))
        row = self.user_rows(importation)[0]
        self.assertEqual(row.action, 'unchanged')
        self.assertEqual(row.warnings, ['usuario_com_outros_dados'])
        existing.refresh_from_db()
        self.assertEqual((existing.name, existing.role.code), ('Elza Souza', 'DOCENTE'))

    def test_inactive_user_is_not_reactivated(self):
        existing = User.objects.register(
            email='elza@ifnmg.edu.br', name='Elza Souza', role=self.teacher_role, username='Elza',
        )
        User.objects.filter(pk=existing.pk).update(is_active=False)
        importation = self.conclude(self.with_users(ELZA))
        self.assertEqual(self.user_rows(importation)[0].warnings, ['usuario_inativo'])
        existing.refresh_from_db()
        self.assertFalse(existing.is_active)

    def test_spreadsheet_without_users_sheet_imports_the_schedule_normally(self):
        importation = self.conclude(workbook_file())
        self.assertEqual(self.user_rows(importation), [])
        self.assertEqual(Teacher.objects.count(), len(TEACHERS))
        self.assertEqual(Space.objects.count(), len(SPACE_NAMES))
        self.assertEqual(Allocation.objects.count(), 4)
        self.assertNotIn('usuário', concluded_message(importation.summary(), importation))

    def test_invalid_email(self):
        importation = self.preview(self.with_users(('elza', 'Elza', 'DOCENTE', 'Elza')))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        self.assertIn(('E-mail', 'invalid'), self.error_codes(self.user_rows(importation)[0]))

    def test_duplicate_email_in_sheet(self):
        importation = self.preview(self.with_users(ELZA, ('ELZA@ifnmg.edu.br', 'Outra', 'DOCENTE', 'Outra')))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        for row in self.user_rows(importation):
            self.assertIn(('E-mail', 'duplicate'), self.error_codes(row))

    def test_duplicate_username_in_sheet(self):
        importation = self.preview(self.with_users(ELZA, ('outra@ifnmg.edu.br', 'Outra', 'DOCENTE', 'ELZA')))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        for row in self.user_rows(importation):
            self.assertIn(('Nome de usuário', 'duplicate'), self.error_codes(row))

    def test_missing_required_cells(self):
        importation = self.preview(self.with_users(
            ('', 'Elza', 'DOCENTE', 'Elza'),
            ('viviane@ifnmg.edu.br', '', 'DOCENTE', 'Viviane'),
            ('erica@ifnmg.edu.br', 'Érica', '', 'Érica'),
            ('ana@ifnmg.edu.br', 'Ana', 'DOCENTE', ''),
        ))
        rows = self.user_rows(importation)
        self.assertEqual(self.error_codes(rows[0]), [('E-mail', 'required')])
        self.assertEqual(self.error_codes(rows[1]), [('Nome', 'required')])
        self.assertEqual(self.error_codes(rows[2]), [('Perfil', 'required')])
        self.assertEqual((rows[3].errors, rows[3].action), ([], 'create'))

    def test_username_taken_by_other_user(self):
        User.objects.register(email='outra@ifnmg.edu.br', name='Outra', role=self.teacher_role, username='Elza')
        importation = self.preview(self.with_users(ELZA))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        self.assertIn(('Nome de usuário', 'username_taken'), self.error_codes(self.user_rows(importation)[0]))

    def test_without_user_create_permission(self):
        self.executor = space_manager('semcadastro@ifnmg.edu.br')
        importation = self.preview(self.with_users(ELZA))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        self.assertIn((None, 'permission_denied'), self.error_codes(self.user_rows(importation)[0]))

    def test_database_error_during_user_creation(self):
        with patch.object(User.objects, 'register', side_effect=DatabaseError):
            importation = self.preview(self.with_users(ELZA))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        self.assertEqual(self.error_codes(self.user_rows(importation)[0]), [(None, 'database')])
        self.assertFalse(self.imported_users().exists())

    def test_user_error_blocks_whole_import(self):
        importation = self.preview(self.with_users(ELZA, ('viviane', 'Viviane', 'DOCENTE', 'Viviane')))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        with self.assertRaises(Exception):
            confirm(importation, self.definition, self.executor)
        self.assertFalse(self.imported_users().exists())
        self.assertFalse(Allocation.objects.exists())

    def test_preview_does_not_persist_users(self):
        importation = self.preview(self.with_users(ELZA, VIVIANE))
        refresh(importation, self.definition, self.executor)
        self.assertFalse(self.imported_users().exists())
        self.assertFalse(self.password_links().exists())
        self.assertFalse(self.user_audit().exists())
        self.assertEqual(User.objects.count(), 1)

    def test_confirmation_persists_users(self):
        importation = self.conclude(self.with_users(ELZA, VIVIANE))
        self.assertEqual(importation.status, Import.Status.CONCLUDED)
        self.assertEqual(self.imported_users().count(), 2)
        self.assertEqual(self.user_audit().count(), 2)
        message = concluded_message(importation.summary(), importation)
        self.assertIn('2 usuários cadastrados', message)
        self.assertTrue(message.endswith('O link para definir a senha será enviado aos usuários cadastrados.'))
        record = AuditLog.objects.get(resource=catalog.SPACE_ALLOCATION_IMPORT, action=catalog.IMPORT)
        self.assertEqual(record.metadata['usuarios_cadastrados'], 2)

    def test_divergence_rolls_back_users(self):
        importation = self.preview(self.with_users(ELZA, VIVIANE))
        User.objects.register(email='elza@ifnmg.edu.br', name='Elza Souza', role=self.teacher_role, username='Elza')
        self.assertFalse(confirm(importation, self.definition, self.executor))
        self.assertEqual(importation.revalidation_count, 1)
        self.assertEqual(importation.status, Import.Status.VALIDATED)
        self.assertFalse(User.objects.filter(email='viviane@ifnmg.edu.br').exists())
        self.assertFalse(Allocation.objects.exists())

    def test_unexpected_error_after_users_rolls_back_everything(self):
        importation = self.preview(self.with_users(ELZA))
        with patch.object(Allocation.objects, 'apply', side_effect=RuntimeError):
            self.assertFalse(confirm(importation, self.definition, self.executor))
        self.assertEqual((importation.status, importation.operation), (Import.Status.VALIDATED, Import.Operation.CONFIRMATION))
        self.assertFalse(self.imported_users().exists())
        self.assertFalse(self.password_links().exists())
        self.assertFalse(Allocation.objects.exists())
        self.assertFalse(Space.objects.exists())
        self.assertFalse(AuditLog.objects.filter(resource=catalog.SPACE).exists())

    def test_confirmer_without_permission(self):
        importation = self.preview(self.with_users(ELZA))
        confirmer = space_manager('semcadastro@ifnmg.edu.br')
        self.assertFalse(confirm(importation, self.definition, confirmer))
        self.assertEqual(importation.status, Import.Status.WITH_ERRORS)
        self.assertFalse(User.objects.filter(email='elza@ifnmg.edu.br').exists())
        self.assertFalse(Allocation.objects.exists())


class SuggestedSpreadsheetTests(ImportTestCase):
    """Planilha sugerida, com a aba `usuários - sistema` (Change importar-usuarios-com-o-ensalamento)."""

    def setUp(self):
        self.executor = space_manager(extra_permissions=[permissions.USER_CREATE])
        self.definition = SpaceAllocationImport()

    def test_users_sheet_is_read_and_every_user_is_planned(self):
        uploaded = SimpleUploadedFile(
            'planilha_example_sugerida.xlsx', SUGGESTED_SPREADSHEET.read_bytes(), content_type=CONTENT_TYPE,
        )
        importation = self.preview(uploaded)
        rows = list(importation.rows.filter(values__Tipo='USUARIO'))
        self.assertEqual(len(rows), 129)
        self.assertEqual([row.errors for row in rows if row.errors], [])
        self.assertEqual({row.action for row in rows}, {'create'})
        self.assertEqual(User.objects.count(), 1)
