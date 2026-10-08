"""Leitura de `sources/planilha_example.xlsx`, com números que acusam mudança de comportamento diante do arquivo real."""

import collections
from unittest.mock import patch

from django.conf import settings
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from imports.reader import FileRejected

from schedule.activities import Catalog, resolve
from schedule.imports import (
    ACTIVITY_COLUMN,
    CELL_COLUMN,
    COURSE_COLUMN,
    DAY_COLUMN,
    GROUP_COLUMN,
    KIND_COLUMN,
    SLOT_COLUMN,
    SpaceAllocationImport,
    sheet_of,
)
from schedule.sheets import (
    ACTIVITY_KIND,
    CATALOG_SHEET,
    SECONDARY_SHEET,
    SHEET_LAYOUTS,
    SPACE_KIND,
    TEACHER_KIND,
    UNDERGRADUATE_SHEET,
    USER_COLUMNS,
    USER_KIND,
    USERS_SHEET,
    catalog_entries,
    group_columns,
    header_rows,
    parse_weekday,
    user_entries,
)

SPREADSHEET = settings.BASE_DIR / 'sources' / 'planilha_example.xlsx'
IGNORED_SHEETS = (
    'Ensalamento - Geral',
    'Horário - Professor',
    'Horário - Professor - 1',
    'Horário - Professor - Separado',
    'Eng. Alimentos',
    'BSI',
    'Lic. em Química',
)


def workbook():
    return load_workbook(SPREADSHEET, data_only=True)


class SpreadsheetTestCase(SimpleTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.definition = SpaceAllocationImport()
        cls.book = workbook()
        cls.rows = list(cls.definition.read_cells(cls.book))
        cls.catalog_rows = [row for row in cls.rows if row[KIND_COLUMN] in (SPACE_KIND, TEACHER_KIND)]
        cls.activity_rows = [row for row in cls.rows if row[KIND_COLUMN] == ACTIVITY_KIND]

    def catalog(self):
        return Catalog.build(
            [self.definition.catalog_name(row) for row in self.catalog_rows if row[KIND_COLUMN] == TEACHER_KIND],
            [self.definition.catalog_name(row) for row in self.catalog_rows if row[KIND_COLUMN] == SPACE_KIND],
        )


class StructureTests(SpreadsheetTestCase):
    def test_portrait_of_the_real_file(self):
        """Números da análise do arquivo real, conferidos juntos para acusar mudança na leitura."""
        self.assertEqual(header_rows(self.book[UNDERGRADUATE_SHEET]), [2, 23, 44, 65, 86, 107])
        self.assertEqual(header_rows(self.book[SECONDARY_SHEET]), [2, 17, 32, 47, 62, 77])
        entries = catalog_entries(self.book)
        teachers = [entry for entry in entries if entry['kind'] == TEACHER_KIND]
        spaces = [entry for entry in entries if entry['kind'] == SPACE_KIND]
        self.assertEqual((len(teachers), len(spaces)), (124, 82))
        self.assertEqual((teachers[0]['value'], teachers[0]['reference']), ('Adailton', f'{CATALOG_SHEET}!A2'))
        self.assertEqual((spaces[0]['value'], spaces[0]['reference']), ('(1/10)', f'{CATALOG_SHEET}!B2'))
        undergraduate = {row[SLOT_COLUMN] for row in self.activity_rows if sheet_of(row) == UNDERGRADUATE_SHEET}
        secondary = {row[SLOT_COLUMN] for row in self.activity_rows if sheet_of(row) == SECONDARY_SHEET}
        self.assertEqual((len(undergraduate), len(secondary)), (15, 9))
        self.assertTrue(secondary < undergraduate)
        self.assertFalse({'INTERVALO', 'ALMOÇO'} & (undergraduate | secondary))
        for title, expected in ((UNDERGRADUATE_SHEET, 39), (SECONDARY_SHEET, 20)):
            layout = next(item for item in SHEET_LAYOUTS if item.title == title)
            sheet = self.book[title]
            columns = group_columns(sheet, layout, header_rows(sheet)[0])
            self.assertEqual(len({column.header_text for column in columns.values()}), expected, title)
        with_activity = {row[GROUP_COLUMN] for row in self.activity_rows if sheet_of(row) == UNDERGRADUATE_SHEET}
        self.assertEqual(len(with_activity), 36)
        origins = [row[CELL_COLUMN] for row in self.activity_rows]
        self.assertIn(f'{UNDERGRADUATE_SHEET}!F69#1', origins)
        self.assertIn(f'{UNDERGRADUATE_SHEET}!F69#2', origins)
        self.assertEqual(len([origin for origin in origins if '#' in origin]), 45)

    def test_space_names_lose_the_parentheses_of_the_list(self):
        names = [self.definition.catalog_name(row) for row in self.catalog_rows if row[KIND_COLUMN] == SPACE_KIND]
        self.assertIn('1/10', names)
        self.assertIn('Lab. de Madeira', names)
        self.assertIn('Lab. Reprodução Animal', names)
        self.assertNotIn('(1/10)', names)

    def test_catalog_rows_come_before_the_activity_rows(self):
        kinds = [row[KIND_COLUMN] for row in self.rows]
        self.assertEqual(set(kinds[:206]), {SPACE_KIND, TEACHER_KIND})
        self.assertEqual(set(kinds[206:]), {ACTIVITY_KIND})

    def test_row_count_of_the_real_file(self):
        self.assertEqual(len(self.catalog_rows), 206)
        self.assertEqual(len(self.activity_rows), 1500)
        self.assertEqual(len(self.rows), 1706)

    def test_calculated_sheets_are_not_read(self):
        for title in IGNORED_SHEETS:
            self.assertIn(title, self.book.sheetnames)
        used = {sheet_of(row) for row in self.rows}
        self.assertEqual(used, {UNDERGRADUATE_SHEET, SECONDARY_SHEET, CATALOG_SHEET})

    def test_the_five_weekdays_of_each_sheet(self):
        for title in (UNDERGRADUATE_SHEET, SECONDARY_SHEET):
            days = {
                parse_weekday(row[DAY_COLUMN]) for row in self.activity_rows if sheet_of(row) == title
            }
            self.assertEqual(days, {'SEGUNDA', 'TERCA', 'QUARTA', 'QUINTA', 'SEXTA'}, title)

    def test_column_without_header_in_one_block_keeps_its_group(self):
        physics = [
            row for row in self.activity_rows
            if sheet_of(row) == UNDERGRADUATE_SHEET and row[GROUP_COLUMN].startswith('4º PERÍODO\nLic. Física')
        ]
        friday = [row for row in physics if parse_weekday(row[DAY_COLUMN]) == 'SEXTA']
        self.assertTrue(friday)
        self.assertEqual({row[GROUP_COLUMN] for row in friday}, {row[GROUP_COLUMN] for row in physics})

    def test_course_title_is_read_except_for_food_engineering(self):
        food = [
            row for row in self.activity_rows
            if row[GROUP_COLUMN].endswith('Eng. Alimentos') or row[GROUP_COLUMN].endswith('Eng. Alimento')
        ]
        self.assertTrue(food)
        self.assertEqual({row[COURSE_COLUMN] for row in food}, {''})
        forest = [row for row in self.activity_rows if row[GROUP_COLUMN].endswith('Eng. Florestal')]
        self.assertEqual({row[COURSE_COLUMN] for row in forest}, {'BACHARELADO EM ENGENHARIA FLORESTAL'})


class ActivityReadingTests(SpreadsheetTestCase):
    def setUp(self):
        self.file_catalog = self.catalog()

    def activity(self, origin):
        row = next(row for row in self.activity_rows if row[CELL_COLUMN] == origin)
        return resolve(row[ACTIVITY_COLUMN], self.file_catalog)

    def test_complete_activity(self):
        activity = self.activity(f'{UNDERGRADUATE_SHEET}!D6')
        self.assertEqual(activity.subject_name, 'CÁLCULO III')
        self.assertEqual(activity.teacher_names, ['Viviane'])
        self.assertEqual(activity.space_name, '2/6')
        self.assertEqual(activity.problems, [])

    def test_subject_with_its_own_parentheses(self):
        activity = self.activity(f'{UNDERGRADUATE_SHEET}!M3')
        self.assertEqual(activity.subject_name, 'Gestão Ambiental (optativa)')
        self.assertEqual(activity.teacher_names, ['Élcio José'])
        self.assertEqual(activity.space_name, '3/10')
        self.assertEqual(activity.problems, [])

    def test_activity_without_space(self):
        activity = self.activity(f'{SECONDARY_SHEET}!F3')
        self.assertEqual(activity.subject_name, 'Ed. Física')
        self.assertEqual(activity.teacher_names, ['Gera Jr.'])
        self.assertIsNone(activity.space_name)
        self.assertEqual(activity.problems, [])

    def test_several_teachers_separated_by_slash(self):
        activity = self.activity(f'{UNDERGRADUATE_SHEET}!V10')
        self.assertEqual(activity.subject_name, 'UCE 2')
        self.assertEqual(
            activity.teacher_names, ['Eduardo Garrido', 'Antônio', 'Lídia'],
        )
        self.assertEqual(activity.space_name, 'HV 3')
        self.assertEqual([code for code, _ in activity.problems], ['docente_desconhecido'])

    def test_several_teachers_separated_by_e(self):
        activity = self.activity(f'{UNDERGRADUATE_SHEET}!AM21')
        self.assertEqual(activity.subject_name, 'Estatística Básica')
        self.assertEqual(activity.teacher_names, ['Arilson'])
        self.assertEqual(activity.space_name, '1/16')
        self.assertIn('Agnaldo', activity.problems[0][1])

    def test_line_break_inside_the_cell(self):
        activity = self.activity(f'{SECONDARY_SHEET}!C3')
        self.assertEqual(activity.subject_name, 'Química')
        self.assertEqual(activity.teacher_names, ['Juliana Diniz'])
        self.assertEqual(activity.space_name, 'Mini 1')

    def test_space_with_an_inner_space_in_the_name(self):
        activity = self.activity(f'{UNDERGRADUATE_SHEET}!U10')
        self.assertEqual(activity.space_name, 'HV 2')
        self.assertEqual(activity.problems, [])

    def test_two_activities_in_one_cell(self):
        first = self.activity(f'{UNDERGRADUATE_SHEET}!F69#1')
        second = self.activity(f'{UNDERGRADUATE_SHEET}!F69#2')
        self.assertEqual(first.subject_name, 'Óleos e Gorduras (OPTATIVA I)')
        self.assertEqual((first.teacher_names, first.space_name), (['Bruna Porto'], 'Sl. Fenômenos'))
        self.assertEqual(second.subject_name, 'Massas e Panificação (OPTATIVA II)')
        self.assertEqual((second.teacher_names, second.space_name), (['Márcia Antunes'], 'Sl. Oficina'))
        self.assertEqual(first.problems + second.problems, [])

    def test_slash_inside_parentheses_does_not_split(self):
        rows = [row for row in self.activity_rows if row[CELL_COLUMN] == f'{UNDERGRADUATE_SHEET}!V3']
        self.assertEqual(len(rows), 1)
        self.assertIn('Eduardo Garrido/ Thaís/ Antônio', rows[0][ACTIVITY_COLUMN])

    def test_space_outside_the_catalog_is_read_as_cited_only_in_the_cells(self):
        activity = self.activity(f'{SECONDARY_SHEET}!J33')
        self.assertEqual(activity.problems, [])
        self.assertEqual((activity.space_name, activity.cell_space_name), (None, 'Zoo I'))

    def test_space_with_a_different_spelling_is_not_matched(self):
        activity = self.activity(f'{UNDERGRADUATE_SHEET}!T25')
        self.assertEqual(activity.problems, [])
        self.assertEqual((activity.space_name, activity.cell_space_name), (None, 'HV1'))
        self.assertEqual(activity.teacher_names, ['Valéria'])

    def test_empty_parentheses_in_place_of_the_space_is_a_problem(self):
        activity = resolve('Física (Elza) ()', Catalog.build(['Elza'], ['1/1']))
        self.assertEqual([code for code, _ in activity.problems], ['sala_vazia'])

    def test_activity_without_subject_is_a_problem(self):
        activity = self.activity(f'{UNDERGRADUATE_SHEET}!AQ64#3')
        self.assertEqual([code for code, _ in activity.problems], ['sem_disciplina'])

    def test_valid_activities_of_a_cell_with_a_problem_are_read(self):
        first = self.activity(f'{UNDERGRADUATE_SHEET}!AQ64#1')
        second = self.activity(f'{UNDERGRADUATE_SHEET}!AQ64#2')
        self.assertEqual(first.subject_name, 'Educação e Diversidade')
        self.assertEqual(first.problems, [])
        self.assertEqual(second.subject_name, 'TCC I - Química')
        self.assertEqual(second.problems, [])

    def test_prefix_is_never_matched(self):
        catalog = Catalog.build(['Gera Jr.', 'Gera Magela'], ['1/1'])
        activity = resolve('Física (Gera) (1/1)', catalog)
        self.assertEqual([code for code, _ in activity.problems], ['docente_desconhecido'])
        self.assertEqual(activity.teacher_names, [])

    def test_accents_are_not_removed_in_the_match(self):
        catalog = Catalog.build(['Érica'], ['1/1'])
        activity = resolve('Biologia (Erica) (1/1)', catalog)
        self.assertEqual([code for code, _ in activity.problems], ['docente_desconhecido'])

    def test_case_is_ignored_in_the_match(self):
        catalog = Catalog.build(['Érica'], ['Mini 1'])
        activity = resolve('Biologia (ÉRICA) (MINI 1)', catalog)
        self.assertEqual(activity.problems, [])
        self.assertEqual((activity.teacher_names, activity.space_name), (['Érica'], 'Mini 1'))

    def test_cell_without_parentheses_is_a_problem(self):
        activity = resolve('Somente um texto', Catalog.build(['Elza'], ['1/1']))
        self.assertEqual([code for code, _ in activity.problems], ['sem_parenteses'])


class RealFileProblemTests(SpreadsheetTestCase):
    def test_how_many_activities_the_real_file_rejects(self):
        file_catalog = self.catalog()
        problems = [resolve(row[ACTIVITY_COLUMN], file_catalog).problems for row in self.activity_rows]
        with_problem = [found for found in problems if found]
        codes = sorted({code for found in with_problem for code, _ in found})
        self.assertEqual(len(with_problem), 121)
        self.assertEqual(codes, ['docente_desconhecido', 'sem_disciplina'])
        by_code = collections.Counter(code for found in with_problem for code, _ in found)
        self.assertEqual(by_code, {'docente_desconhecido': 120, 'sem_disciplina': 1})


class UserSheetReadingTests(SimpleTestCase):
    """Leitura da aba opcional `usuários - sistema`."""

    def book(self, rows=None, title=USERS_SHEET):
        book = Workbook()
        book.remove(book.active)
        catalog_sheet = book.create_sheet(CATALOG_SHEET)
        catalog_sheet['A1'] = 'Disciplina (Professor)'
        catalog_sheet['B1'] = 'Salas Disponíveis'
        catalog_sheet['A2'] = 'Elza'
        catalog_sheet['B2'] = '(1/13)'
        if rows is not None:
            sheet = book.create_sheet(title)
            for row in rows:
                sheet.append(list(row))
        return book

    def test_missing_sheet_gives_no_users(self):
        self.assertEqual(user_entries(self.book()), [])

    def test_one_entry_per_data_row_with_raw_values(self):
        entries = user_entries(self.book([
            USER_COLUMNS,
            (' elza@ifnmg.edu.br ', 'Elza', 'DOCENTE', 'Elza'),
            (None, None, None, None),
            ('viviane@ifnmg.edu.br', 'Viviane', 'DOCENTE', None),
        ]))
        self.assertEqual([entry['reference'] for entry in entries], ['usuários - sistema!A2', 'usuários - sistema!A4'])
        self.assertEqual(entries[0]['values']['E-mail'], ' elza@ifnmg.edu.br ')
        self.assertEqual(entries[1]['values']['Nome de usuário'], '')

    def test_sheet_title_without_case_and_accents(self):
        book = self.book([USER_COLUMNS, ('elza@ifnmg.edu.br', 'Elza', 'DOCENTE', 'Elza')], title='Usuarios - Sistema')
        self.assertEqual(user_entries(book)[0]['reference'], 'Usuarios - Sistema!A2')

    def test_header_without_case_and_accents_and_extra_column(self):
        entries = user_entries(self.book([
            ('e-mail', 'NOME', 'Observação', 'perfil', 'nome de usuario'),
            ('elza@ifnmg.edu.br', 'Elza', 'qualquer', 'DOCENTE', 'Elza'),
        ]))
        self.assertEqual(
            entries[0]['values'],
            {'E-mail': 'elza@ifnmg.edu.br', 'Nome': 'Elza', 'Perfil': 'DOCENTE', 'Nome de usuário': 'Elza'},
        )

    def test_missing_column_rejects_the_file(self):
        with self.assertRaises(FileRejected) as raised:
            user_entries(self.book([('E-mail', 'Nome', 'Nome de usuário'), ('a@ifnmg.edu.br', 'A', 'A')]))
        self.assertIn('"Perfil"', raised.exception.message)
        self.assertIn('"usuários - sistema"', raised.exception.message)

    def test_repeated_column_rejects_the_file(self):
        with self.assertRaises(FileRejected) as raised:
            user_entries(self.book([(*USER_COLUMNS, 'Perfil')]))
        self.assertIn('"Perfil" aparece mais de uma vez', raised.exception.message)

    def test_empty_sheet_and_header_only_give_no_users(self):
        self.assertEqual(user_entries(self.book([])), [])
        self.assertEqual(user_entries(self.book([USER_COLUMNS])), [])

    def test_user_rows_come_before_catalog_rows(self):
        book = self.book([USER_COLUMNS, ('elza@ifnmg.edu.br', 'Elza', 'DOCENTE', 'Elza')])
        with patch('schedule.imports.SHEET_LAYOUTS', ()):
            rows = list(SpaceAllocationImport().read_cells(book))
        self.assertEqual([row[KIND_COLUMN] for row in rows], [USER_KIND, SPACE_KIND, TEACHER_KIND])
